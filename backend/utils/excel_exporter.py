"""
Excel导出工具类

提供带有自动列宽调整和美化功能的Excel导出能力
"""

import io
from typing import Dict, List, Optional, Union
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


class ExcelExporter:
    """Excel导出工具类

    特性:
    - 自动列宽调整（支持最大宽度限制）
    - 表头美化（浅蓝色背景）
    - 行间隔美化（浅灰色背景）
    - 边框美化
    """

    # 默认配置
    DEFAULT_MAX_COLUMN_WIDTH = 50  # 最大列宽
    DEFAULT_MIN_COLUMN_WIDTH = 10  # 最小列宽

    # 颜色配置（克制的浅色系）
    HEADER_FILL_COLOR = "D6E4F5"  # 浅蓝色（表头）
    ALTERNATE_ROW_COLOR = "F2F2F2"  # 浅灰色（间隔行）

    def __init__(
        self,
        max_column_width: int = DEFAULT_MAX_COLUMN_WIDTH,
        min_column_width: int = DEFAULT_MIN_COLUMN_WIDTH,
        enable_styling: bool = True
    ):
        """初始化Excel导出器

        Args:
            max_column_width: 最大列宽
            min_column_width: 最小列宽
            enable_styling: 是否启用样式美化
        """
        self.max_column_width = max_column_width
        self.min_column_width = min_column_width
        self.enable_styling = enable_styling

    def _calculate_column_width(self, column_data: pd.Series, column_name: str) -> float:
        """计算列的最佳宽度

        Args:
            column_data: 列数据
            column_name: 列名

        Returns:
            计算后的列宽
        """
        # 计算列名的长度（中文字符占2个宽度）
        def calc_string_width(s: str) -> int:
            """计算字符串显示宽度（中文字符算2个宽度）"""
            width = 0
            for char in str(s):
                # 简单判断：中文字符范围
                if '\u4e00' <= char <= '\u9fff':
                    width += 2
                else:
                    width += 1
            return width

        # 列名宽度
        header_width = calc_string_width(column_name)

        # 数据内容最大宽度
        if len(column_data) > 0:
            # 采样前100行以提升性能
            sample_data = column_data.head(100)
            max_content_width = sample_data.astype(str).apply(calc_string_width).max()
        else:
            max_content_width = 0

        # 取较大值，并加上一些padding
        optimal_width = max(header_width, max_content_width) + 2

        # 限制在最小和最大宽度之间
        return max(
            self.min_column_width,
            min(optimal_width, self.max_column_width)
        )

    def _apply_header_style(self, worksheet: Worksheet, header_row: int = 1):
        """应用表头样式

        Args:
            worksheet: 工作表对象
            header_row: 表头行号（1-based）
        """
        if not self.enable_styling:
            return

        # 表头样式
        header_font = Font(name='Arial', size=11, bold=True, color="000000")
        header_fill = PatternFill(start_color=self.HEADER_FILL_COLOR,
                                  end_color=self.HEADER_FILL_COLOR,
                                  fill_type="solid")
        header_alignment = Alignment(horizontal='center', vertical='center')

        # 应用到表头行的所有单元格
        for cell in worksheet[header_row]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

    def _apply_alternate_row_colors(self, worksheet: Worksheet, start_row: int = 2, end_row: Optional[int] = None):
        """应用行间隔颜色

        Args:
            worksheet: 工作表对象
            start_row: 起始行号（1-based）
            end_row: 结束行号（1-based），None表示到最后一行
        """
        if not self.enable_styling:
            return

        if end_row is None:
            end_row = worksheet.max_row

        # 间隔行填充色
        alternate_fill = PatternFill(start_color=self.ALTERNATE_ROW_COLOR,
                                     end_color=self.ALTERNATE_ROW_COLOR,
                                     fill_type="solid")

        # 应用到偶数行
        for row_idx in range(start_row, end_row + 1):
            if row_idx % 2 == 0:  # 偶数行
                for cell in worksheet[row_idx]:
                    cell.fill = alternate_fill

    def _apply_borders(self, worksheet: Worksheet):
        """应用边框

        Args:
            worksheet: 工作表对象
        """
        if not self.enable_styling:
            return

        # 细边框样式
        thin_border = Border(
            left=Side(style='thin', color='D3D3D3'),
            right=Side(style='thin', color='D3D3D3'),
            top=Side(style='thin', color='D3D3D3'),
            bottom=Side(style='thin', color='D3D3D3')
        )

        # 应用到所有有数据的单元格
        for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row,
                                       min_col=1, max_col=worksheet.max_column):
            for cell in row:
                cell.border = thin_border

    def _auto_adjust_column_widths(self, worksheet: Worksheet, df: pd.DataFrame):
        """自动调整列宽

        Args:
            worksheet: 工作表对象
            df: DataFrame数据
        """
        for idx, column in enumerate(df.columns, start=1):
            column_letter = get_column_letter(idx)
            column_data = df[column]

            # 计算最佳宽度
            optimal_width = self._calculate_column_width(column_data, column)

            # 应用宽度
            worksheet.column_dimensions[column_letter].width = optimal_width

    def style_worksheet(self, worksheet: Worksheet, df: pd.DataFrame):
        """对工作表应用完整样式

        Args:
            worksheet: 工作表对象
            df: DataFrame数据
        """
        # 自动调整列宽
        self._auto_adjust_column_widths(worksheet, df)

        # 应用表头样式
        self._apply_header_style(worksheet, header_row=1)

        # 应用行间隔颜色（从第2行开始，跳过表头）
        self._apply_alternate_row_colors(worksheet, start_row=2)

        # 应用边框
        self._apply_borders(worksheet)

    def export_to_bytes(self, data_sheets: Dict[str, pd.DataFrame]) -> io.BytesIO:
        """导出多个工作表到BytesIO对象

        Args:
            data_sheets: 工作表字典，key为sheet名称，value为DataFrame

        Returns:
            包含Excel内容的BytesIO对象
        """
        output = io.BytesIO()

        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            # 写入所有工作表
            for sheet_name, df in data_sheets.items():
                df.to_excel(writer, sheet_name=sheet_name, index=False)

            # 应用样式到每个工作表
            for sheet_name, df in data_sheets.items():
                worksheet = writer.sheets[sheet_name]
                self.style_worksheet(worksheet, df)

        output.seek(0)
        return output

    def export_single_sheet(self, df: pd.DataFrame, sheet_name: str = "Sheet1") -> io.BytesIO:
        """导出单个工作表到BytesIO对象

        Args:
            df: DataFrame数据
            sheet_name: 工作表名称

        Returns:
            包含Excel内容的BytesIO对象
        """
        return self.export_to_bytes({sheet_name: df})


def create_excel_exporter(
    max_column_width: int = ExcelExporter.DEFAULT_MAX_COLUMN_WIDTH,
    min_column_width: int = ExcelExporter.DEFAULT_MIN_COLUMN_WIDTH,
    enable_styling: bool = True
) -> ExcelExporter:
    """创建Excel导出器实例（工厂函数）

    Args:
        max_column_width: 最大列宽
        min_column_width: 最小列宽
        enable_styling: 是否启用样式美化

    Returns:
        ExcelExporter实例
    """
    return ExcelExporter(
        max_column_width=max_column_width,
        min_column_width=min_column_width,
        enable_styling=enable_styling
    )
